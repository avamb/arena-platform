from __future__ import annotations

import json
import os
import tempfile
import unittest
from openpyxl import Workbook
from pathlib import Path

from m1_product.cli import _run_dir, fixture_acceptance, overall, pack, pack_venue


def _aribau_synthetic(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Aribau,Sala 5 Aforo ok"
    ws["A1"] = "PARTERRE"
    for col, v in zip("ABCDE", [1, 1, 2, 3, 1]):
        ws[f"{col}2"] = v
    wb.save(path)


class M1Tests(unittest.TestCase):
    def test_missing_inputs_never_pass(self):
        with tempfile.TemporaryDirectory() as tmp:
            summary = pack(None, None, Path(tmp))
            self.assertEqual(summary["overall_status"], "INCOMPLETE")
            self.assertEqual(summary["fixture_acceptance"]["status"], "NOT_VERIFIED")
            self.assertTrue(all(r["status"] == "MISSING_INPUT" for r in summary["venues"]))
            self.assertTrue((Path(tmp) / "validation-summary.json").is_file())

    def test_non_xlsx_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            rec = pack_venue("hunter", Path(tmp) / "ref.svg", Path(tmp))
            self.assertEqual(rec["status"], "FAIL")

    def test_capacity_mismatch_is_not_pass(self):
        with tempfile.TemporaryDirectory() as tmp:
            x = Path(tmp) / "a.xlsx"
            _aribau_synthetic(x)
            rec = pack_venue("aribau", x, Path(tmp), frozen=Path(tmp) / "none.json")
            self.assertEqual(rec["seat_count"], 3)
            self.assertEqual(rec["status"], "FAIL")
            self.assertEqual(rec["frozen_regression"]["status"], "NOT_AVAILABLE")

    def test_synthetic_pass_with_overridden_target(self):
        with tempfile.TemporaryDirectory() as tmp:
            x = Path(tmp) / "a.xlsx"
            _aribau_synthetic(x)
            rec = pack_venue("aribau", x, Path(tmp), target=3, frozen=Path(tmp) / "none.json")
            self.assertIn(rec["status"], {"PASS", "QUARANTINE"}, rec)
            if rec["status"] == "PASS":
                self.assertEqual(rec["render_validation"]["seat_count"], 3)
            self.assertEqual(len(rec["input_sha256"]), 64)

    def test_frozen_baseline_is_mandatory_for_product_pass(self):
        records = [
            {"venue": "hunter", "status": "PASS", "frozen_regression": {"status": "NOT_AVAILABLE"}},
            {"venue": "aribau", "status": "PASS", "frozen_regression": {"status": "NOT_AVAILABLE"}},
        ]
        gate = fixture_acceptance(records)
        self.assertEqual(gate["status"], "NOT_VERIFIED")
        self.assertTrue(gate["required"])

    def test_each_cli_run_gets_distinct_output_directory(self):
        with tempfile.TemporaryDirectory() as tmp:
            first, first_id = _run_dir(Path(tmp), "run-one")
            (first / "hunter.json").write_text("stale", encoding="utf-8")
            second, second_id = _run_dir(Path(tmp), "run-two")
            self.assertNotEqual(first, second)
            self.assertNotEqual(first_id, second_id)
            self.assertTrue((first / "hunter.json").is_file())
            self.assertFalse((second / "hunter.json").exists())
            with self.assertRaises(FileExistsError):
                _run_dir(Path(tmp), "run-two")

    def test_partial_run_cannot_reuse_previous_hunter_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first, _ = _run_dir(root, "full-run")
            (first / "hunter.json").write_text("old hunter", encoding="utf-8")
            second, _ = _run_dir(root, "aribau-only-run")
            aribau = root / "aribau.xlsx"
            _aribau_synthetic(aribau)
            summary = pack(None, aribau, second)
            self.assertEqual(summary["overall_status"], "INCOMPLETE")
            self.assertFalse((second / "hunter.json").exists())
            self.assertTrue((first / "hunter.json").exists())

    def test_overall_logic(self):
        self.assertEqual(overall([{"status": "PASS"}, {"status": "MISSING_INPUT"}]), "INCOMPLETE")
        self.assertEqual(overall([{"status": "PASS"}, {"status": "QUARANTINE"}]), "QUARANTINE")
        self.assertEqual(overall([{"status": "PASS"}, {"status": "FAIL"}]), "FAIL")
        self.assertEqual(overall([{"status": "PASS"}, {"status": "PASS"}]), "PASS")

    @unittest.skipUnless(os.environ.get("M1_HUNTER_XLSX") and os.environ.get("M1_ARIBAU_XLSX"),
                         "real fixtures unavailable: set M1_HUNTER_XLSX and M1_ARIBAU_XLSX")
    def test_real_fixtures(self):
        with tempfile.TemporaryDirectory() as tmp:
            s = pack(Path(os.environ["M1_HUNTER_XLSX"]), Path(os.environ["M1_ARIBAU_XLSX"]), Path(tmp))
            seats = {r["venue"]: r["seat_count"] for r in s["venues"]}
            self.assertEqual(seats, {"hunter": 907, "aribau": 1166})
            if os.environ.get("M1_HUNTER_BASELINE") and os.environ.get("M1_ARIBAU_BASELINE"):
                self.assertEqual(s["overall_status"], "PASS")
            else:
                self.assertEqual(s["overall_status"], "NOT_VERIFIED")


if __name__ == "__main__":
    unittest.main()

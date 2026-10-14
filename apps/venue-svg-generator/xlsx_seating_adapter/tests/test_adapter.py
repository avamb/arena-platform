from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from xlsx_seating_adapter.core import extract_workbook, validate_model
from xlsx_seating_adapter.profiles import ARIBAU_PROFILE, HUNTER_PROFILE


ROOT = Path(__file__).resolve().parents[2]
HUNTER_XLSX = ROOT / "ChatExport_2026-10-11" / "files" / "Площадка .xlsx"
ARIBAU_XLSX = ROOT / "ChatExport_2026-10-11" / "files" / "AFORO ARIBAU SALA Barcelona.xlsx"
PHASE7B = ROOT / "benchmark_v1" / "phase7b_hunter" / "output" / "seating-model-source-only.json"
PHASE8C = ROOT / "benchmark_v1" / "phase8c_aribau" / "output" / "aribau-full-source-only.json"


def semantic_index(model):
    return {
        (seat["section_label"], seat["row_label"], seat["seat_number"]): seat["source_coordinate"]["cell"]
        for seat in model["seats"]
        if seat["number_status"] == "explicitly_read"
    }


class AdapterRegressionTests(unittest.TestCase):
    def test_hunter_semantics_match_phase7b(self):
        model = extract_workbook(HUNTER_XLSX, HUNTER_PROFILE)
        frozen = json.loads(PHASE7B.read_text(encoding="utf-8"))
        new_index = {(seat["row_label"], seat["seat_number"]): seat["source_coordinate"]["cell"] for seat in model["seats"]}
        old_index = {(seat["row_label"], seat["seat_number"]): seat["source_coordinate"]["cell"] for seat in frozen["seats"]}
        self.assertEqual(model["validation"]["status"], "PASS")
        self.assertEqual(len(new_index), 907)
        self.assertEqual(new_index, old_index)
        self.assertEqual(model["summary"]["unknown_count"], 0)

    def test_aribau_semantics_match_phase8c(self):
        model = extract_workbook(ARIBAU_XLSX, ARIBAU_PROFILE)
        frozen = json.loads(PHASE8C.read_text(encoding="utf-8"))
        self.assertEqual(model["validation"]["status"], "PASS")
        self.assertEqual(model["summary"]["seat_count"], 1166)
        self.assertEqual(semantic_index(model), semantic_index(frozen))
        self.assertEqual([(item["label"], item["seat_count"]) for item in model["sections"]], [("BALCONY", 266), ("PARTERRE", 900), ("STAGE", 0)])

    def test_merged_headers_service_cells_repeated_labels_and_gaps(self):
        aribau = extract_workbook(ARIBAU_XLSX, ARIBAU_PROFILE)
        hunter = extract_workbook(HUNTER_XLSX, HUNTER_PROFILE)
        self.assertIn("I41:T41", aribau["source"]["sheets"][0]["merged_cells"])
        row21 = next(row for row in aribau["rows"] if row["worksheet_row"] == 21)
        self.assertEqual([item["cell"] for item in row21["gaps"]], ["H21", "U21"])
        reasons = {item["reason"] for item in aribau["rejected_cells"]}
        self.assertIn("section_header_not_seat", reasons)
        self.assertIn("row_label_marker", reasons)
        self.assertTrue(all(len(row["marker_cells"]) == 4 for row in hunter["rows"]))
        self.assertEqual(sum(item["reason"] == "row_label_marker" for item in aribau["rejected_cells"]), 108)

    def _synthetic(self, values, merge=True):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Test"
        sheet["B1"] = "PARTERRE"
        if merge:
            sheet.merge_cells("B1:D1")
        sheet["A2"] = 5
        for column, value in zip((2, 3, 4), values):
            sheet.cell(2, column).value = value
        sheet["E2"] = 5
        handle = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
        handle.close()
        path = Path(handle.name)
        workbook.save(path)
        profile = {"profile_id": "synthetic", "sheet_name": "Test", "section_header_labels": ["PARTERRE"], "row_detection": {"type": "equal_numeric_edge_markers"}, "seat_detection": {"type": "numeric_between_edge_markers"}}
        return path, profile

    def test_unknown_is_quarantined(self):
        path, profile = self._synthetic((1.5, None, 2))
        try:
            model = extract_workbook(path, profile)
            self.assertEqual(model["validation"]["status"], "QUARANTINE")
            self.assertEqual(model["summary"]["unknown_count"], 1)
            self.assertEqual(validate_model(model)["status"], "QUARANTINE")
        finally:
            path.unlink(missing_ok=True)

    def test_duplicate_is_quarantined(self):
        path, profile = self._synthetic((1, 1, 2), merge=False)
        try:
            model = extract_workbook(path, profile)
            self.assertEqual(model["validation"]["status"], "QUARANTINE")
            self.assertEqual(len(model["summary"]["duplicate_identity_keys"]), 1)
        finally:
            path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()

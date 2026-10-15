"""Limited fail-closed XLSX -> Seating JSON implementation.

Supported structures are deliberately small: repeated single-letter row
markers with a configured style allowlist, and sections whose rows have equal
numeric edge markers with numeric seats between them. Reference SVGs are not
accepted as an input to this module.
"""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SCHEMA_VERSION = "seating-json-v1"
ADAPTER_VERSION = "0.1.0"
ALLOWED_REJECTED_REASONS = {"numeric_service_or_marker", "row_label_marker", "section_header_not_seat"}


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _cell_evidence(cell: Any) -> dict[str, Any]:
    fill = cell.fill.fgColor
    return {
        "sheet": cell.parent.title,
        "cell": cell.coordinate,
        "row": cell.row,
        "column": cell.column,
        "column_letter": get_column_letter(cell.column),
        "raw_value": cell.value,
        "data_type": cell.data_type,
        "style_id": cell.style_id,
        "fill_rgb": str(getattr(fill, "rgb", fill)) if fill is not None else None,
    }


def _merged_range_for(sheet: Any, cell: Any) -> str | None:
    for merged in sheet.merged_cells.ranges:
        if cell.coordinate in merged:
            return str(merged)
    return None


def inspect_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        sheets.append({
            "title": sheet.title,
            "used_range": sheet.calculate_dimension(),
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "merged_cells": [str(item) for item in sheet.merged_cells.ranges],
            "formula_cells": [cell.coordinate for row in sheet.iter_rows() for cell in row if cell.data_type == "f"],
            "hidden_rows": [idx for idx, dim in sheet.row_dimensions.items() if dim.hidden],
            "hidden_columns": [idx for idx, dim in sheet.column_dimensions.items() if dim.hidden],
            "nonempty_cell_count": sum(cell.value is not None for row in sheet.iter_rows() for cell in row),
            "numeric_cell_count": sum(_is_number(cell.value) for row in sheet.iter_rows() for cell in row),
        })
    return {"path": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest(), "workbook_sheets": workbook.sheetnames, "sheets": sheets}


def _section_ranges(sheet: Any, labels: set[str]) -> list[dict[str, Any]]:
    headers = []
    for row in range(1, sheet.max_row + 1):
        for cell in sheet[row]:
            if isinstance(cell.value, str) and cell.value.strip().upper() in labels:
                headers.append({"label": cell.value.strip(), "row": row, "evidence": {**_cell_evidence(cell), "merged_range": _merged_range_for(sheet, cell)}})
    return [
        {"section_id": f"section-{i + 1}", "label": header["label"], "header_row": header["row"], "end_row_exclusive": headers[i + 1]["row"] if i + 1 < len(headers) else sheet.max_row + 1, "header": header}
        for i, header in enumerate(headers)
    ]


def _detect_markers(row_cells: list[Any], profile: dict[str, Any]):
    numeric = sorted((cell for cell in row_cells if _is_number(cell.value)), key=lambda cell: cell.column)
    strategy = profile["row_detection"]["type"]
    if strategy == "single_letter_markers":
        allowed = profile["row_detection"]["allowed_labels"]
        markers = [cell for cell in row_cells if isinstance(cell.value, str) and len(cell.value.strip()) == 1 and cell.value.strip().upper() in allowed]
        return ({"label": markers[0].value.strip().upper(), "cells": markers}, numeric, None) if markers else (None, numeric, "no_single_letter_row_marker")
    if strategy == "equal_numeric_edge_markers":
        if len(numeric) < 2:
            return None, numeric, "insufficient_numeric_edge_markers"
        if numeric[0].value != numeric[-1].value:
            return None, numeric, "edge_numeric_markers_disagree"
        return {"label": str(numeric[0].value), "cells": [numeric[0], numeric[-1]]}, numeric, None
    raise ValueError(f"Unsupported row detection strategy: {strategy}")


def _candidate_cells(row_cells: list[Any], markers: dict[str, Any], profile: dict[str, Any]):
    strategy = profile["seat_detection"]["type"]
    numeric = sorted((cell for cell in row_cells if _is_number(cell.value)), key=lambda cell: cell.column)
    if strategy == "style_allowlist":
        allowed = set(profile["seat_detection"]["style_ids"])
        seats = [cell for cell in numeric if cell.style_id in allowed]
        rejected = [{"reason": "numeric_service_or_marker", "evidence": _cell_evidence(cell)} for cell in numeric if cell.style_id not in allowed]
        return seats, [], rejected
    if strategy == "numeric_between_edge_markers":
        left, right = markers["cells"][0], markers["cells"][-1]
        interior = [cell for cell in row_cells if left.column < cell.column < right.column]
        rejected = [{"reason": "non_numeric_interior_cell", "evidence": _cell_evidence(cell)} for cell in interior if cell.value is not None and not _is_number(cell.value)]
        return [cell for cell in interior if _is_number(cell.value)], [cell for cell in interior if cell.value is None], rejected
    raise ValueError(f"Unsupported seat detection strategy: {strategy}")


def _quarantine(inspection: dict[str, Any], profile: dict[str, Any], reason: str) -> dict[str, Any]:
    return {
        "schema_version": SCHEMA_VERSION,
        "adapter_version": ADAPTER_VERSION,
        "source_only": True,
        "reference_svg_used": False,
        "profile_id": profile["profile_id"],
        "source": inspection,
        "sections": [], "rows": [], "seats": [], "gaps": [], "rejected_cells": [{"reason": reason}],
        "summary": {"section_count": 0, "row_count": 0, "seat_count": 0, "explicit_number_count": 0, "unknown_count": 0, "gap_count": 0, "rejected_count": 1, "duplicate_identity_keys": [], "duplicate_seat_keys": [], "position_numbering_used": False},
        "validation": {"status": "QUARANTINE", "issues": [reason], "all_coordinates_present": True},
    }


def extract_workbook(path: Path, profile: dict[str, Any]) -> dict[str, Any]:
    inspection = inspect_workbook(path)
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheet_name = profile["sheet_name"]
    if sheet_name not in workbook.sheetnames:
        return _quarantine(inspection, profile, f"configured_sheet_missing:{sheet_name}")
    sheet = workbook[sheet_name]
    sections = _section_ranges(sheet, set(profile.get("section_header_labels", [])))
    if not sections:
        sections = [{"section_id": "section-1", "label": None, "header_row": sheet.min_row - 1, "end_row_exclusive": sheet.max_row + 1, "header": None}]
    rows, seats, gaps, rejected = [], [], [], []
    for section in sections:
        if section["header"] is not None:
            item = {"reason": "section_header_not_seat", "section_id": section["section_id"], "evidence": section["header"]["evidence"]}
            rejected.append(item)
        for row_number in range(section["header_row"] + 1, section["end_row_exclusive"]):
            row_cells = list(sheet[row_number])
            markers, numeric, marker_error = _detect_markers(row_cells, profile)
            if markers is None:
                if numeric:
                    row_rejected = [{"reason": marker_error or "row_marker_unknown", "evidence": _cell_evidence(cell)} for cell in numeric]
                    rows.append({"geometric_row_id": f"worksheet-row-{row_number}", "worksheet_row": row_number, "official_label": None, "marker_cells": [], "section_id": section["section_id"], "section_label": section["label"], "seat_count": 0, "unknown_count": 0, "gaps": [], "rejected_cells": row_rejected, "seats": []})
                    rejected.extend(row_rejected)
                continue
            candidate_cells, row_gaps, row_rejected = _candidate_cells(row_cells, markers, profile)
            row_rejected.extend({"reason": "row_label_marker", "evidence": _cell_evidence(marker)} for marker in markers["cells"])
            row_seats = []
            for cell in candidate_cells:
                integer = float(cell.value).is_integer()
                number_status = "explicitly_read" if integer else "unknown"
                seat = {
                    "seat_key": f"xlsx-seat-{sheet.title}-{cell.coordinate}",
                    "source_coordinate": _cell_evidence(cell),
                    "computed_geometry": {"grid_x": cell.column, "grid_y": cell.row, "column_letter": get_column_letter(cell.column)},
                    "section_id": section["section_id"], "section_label": section["label"], "geometric_row_id": f"worksheet-row-{row_number}", "row_label": str(markers["label"]),
                    "seat_number": int(cell.value) if integer else None, "number_status": number_status, "confidence": 1.0 if integer else 0.0,
                    "evidence": {"sheet": sheet.title, "source_cell": cell.coordinate, "raw_value": cell.value, "style_id": cell.style_id, "row_marker_cells": [marker.coordinate for marker in markers["cells"]], "position_numbering_used": False, "rule_reconstruction_used": False},
                }
                row_seats.append(seat)
                if not integer:
                    row_rejected.append({"reason": "unknown_numeric_candidate", "evidence": _cell_evidence(cell)})
            row_gap_records = []
            for cell in row_gaps:
                record = {"cell": cell.coordinate, "row": cell.row, "column": cell.column, "reason": "empty_position_inside_row_span", "merged_range": _merged_range_for(sheet, cell)}
                row_gap_records.append(record)
                gaps.append({"section_id": section["section_id"], "geometric_row_id": f"worksheet-row-{row_number}", **record})
            row_record = {"geometric_row_id": f"worksheet-row-{row_number}", "worksheet_row": row_number, "official_label": str(markers["label"]), "marker_cells": [marker.coordinate for marker in markers["cells"]], "section_id": section["section_id"], "section_label": section["label"], "seat_count": len(row_seats), "unknown_count": sum(seat["number_status"] == "unknown" for seat in row_seats), "gaps": row_gap_records, "rejected_cells": row_rejected, "seats": row_seats}
            rows.append(row_record); seats.extend(row_seats); rejected.extend(row_rejected)
    identities = [(seat["section_label"], seat["row_label"], seat["seat_number"]) for seat in seats if seat["number_status"] == "explicitly_read"]
    duplicate_identity = sorted({identity for identity in identities if identities.count(identity) > 1})
    duplicate_keys = sorted({seat["seat_key"] for seat in seats if [item["seat_key"] for item in seats].count(seat["seat_key"]) > 1})
    issues = []
    if duplicate_identity: issues.append("duplicate_identity_keys")
    if duplicate_keys: issues.append("duplicate_seat_keys")
    if any(seat["number_status"] == "unknown" for seat in seats): issues.append("unknown_numbers")
    if any(item["reason"] not in ALLOWED_REJECTED_REASONS for item in rejected): issues.append("unclassified_rejected_cells")
    section_summary = []
    for section in sections:
        section_rows = [row for row in rows if row["section_id"] == section["section_id"]]
        section_summary.append({"section_id": section["section_id"], "label": section["label"], "header_row": section["header_row"], "row_count": len(section_rows), "seat_count": sum(row["seat_count"] for row in section_rows), "unknown_count": sum(row["unknown_count"] for row in section_rows), "gap_count": sum(len(row["gaps"]) for row in section_rows), "rejected_count": sum(len(row["rejected_cells"]) for row in section_rows) + (1 if section["header"] is not None else 0)})
    result = {"schema_version": SCHEMA_VERSION, "adapter_version": ADAPTER_VERSION, "source_only": True, "reference_svg_used": False, "profile_id": profile["profile_id"], "source": inspection, "sections": section_summary, "rows": rows, "seats": seats, "gaps": gaps, "rejected_cells": rejected, "summary": {"section_count": len(section_summary), "row_count": len(rows), "seat_count": len(seats), "explicit_number_count": sum(seat["number_status"] == "explicitly_read" for seat in seats), "unknown_count": sum(seat["number_status"] == "unknown" for seat in seats), "gap_count": len(gaps), "rejected_count": len(rejected), "duplicate_identity_keys": [list(item) for item in duplicate_identity], "duplicate_seat_keys": duplicate_keys, "position_numbering_used": False}, "validation": {"status": "PASS" if not issues else "QUARANTINE", "issues": issues, "all_coordinates_present": all(bool(seat["source_coordinate"]["cell"]) for seat in seats)}}
    return result


def validate_model(model: dict[str, Any]) -> dict[str, Any]:
    issues = []
    if model.get("schema_version") != SCHEMA_VERSION: issues.append("schema_version")
    if not model.get("source_only") or model.get("reference_svg_used"): issues.append("reference_or_source_latch")
    if any(not seat.get("source_coordinate", {}).get("cell") for seat in model.get("seats", [])): issues.append("missing_source_coordinate")
    if model.get("summary", {}).get("duplicate_identity_keys"): issues.append("duplicate_identity_keys")
    if model.get("summary", {}).get("duplicate_seat_keys"): issues.append("duplicate_seat_keys")
    if model.get("summary", {}).get("unknown_count", 0): issues.append("unknown_numbers")
    return {"status": "PASS" if not issues else "QUARANTINE", "issues": issues}


def write_json(model: dict[str, Any], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(model, ensure_ascii=False, indent=2), encoding="utf-8")

"""Builds a small synthetic workbook of the `synthetic-v1` profile shape.

No client data: the layout is invented. It exists so the pipeline, the HTTP service and the
arena import proof (`arena_export/testdata`) can be exercised on a clean checkout.
"""

from __future__ import annotations

from pathlib import Path
from typing import Optional, Sequence

from openpyxl import Workbook

# (section label, rows). A row is (row number, seat values); None is an empty cell (an aisle gap)
# and any other value is written as is (a float like 1.5 makes the adapter quarantine the model).
SyntheticLayout = Sequence[tuple[str, Sequence[tuple[int, Sequence[Optional[float]]]]]]

DEFAULT_LAYOUT: SyntheticLayout = (
    ("PARTERRE", (
        (1, (1, 2, 3, None, 4, 5)),
        (2, (1, 2, 3, 4, 5, 6)),
        (3, (1, 2, 3, 4, 5, 6)),
    )),
    ("BALCONY", (
        (1, (1, 2, 3, 4)),
        (2, (1, 2, 3, 4)),
    )),
)


def build_synthetic_workbook(path: str | Path, layout: SyntheticLayout = DEFAULT_LAYOUT, sheet_name: str = "Seating") -> Path:
    """Write the workbook and return its path. Cells: column A is the left row marker, seats follow,
    the right row marker closes the row; a blank worksheet row separates sections."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    current = 1
    for label, rows in layout:
        sheet.cell(current, 2).value = label
        current += 1
        for row_number, seats in rows:
            sheet.cell(current, 1).value = row_number
            for offset, value in enumerate(seats):
                sheet.cell(current, 2 + offset).value = value
            sheet.cell(current, 2 + len(seats)).value = row_number
            current += 1
        current += 1
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out)
    return out

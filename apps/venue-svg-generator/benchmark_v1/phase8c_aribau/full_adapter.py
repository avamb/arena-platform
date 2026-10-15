"""Bounded full-workbook runner reusing the Phase 8B generic XLSX primitives.

The venue-specific decisions live in the profile dictionaries. This module
contains only the two configured classification strategies needed for the
Hunter regression and Aribau source-only inventory.
"""

from __future__ import annotations

import importlib.util
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ADAPTER_DIR = Path(__file__).resolve().parents[1] / "phase8b_aribau" / "adapter"
_generic_spec = importlib.util.spec_from_file_location("phase8b_generic_xlsx", ADAPTER_DIR / "generic_xlsx.py")
if _generic_spec is None or _generic_spec.loader is None:
    raise ImportError("Cannot load Phase 8B generic XLSX adapter")
_generic = importlib.util.module_from_spec(_generic_spec)
_generic_spec.loader.exec_module(_generic)
_cell_evidence = _generic._cell_evidence
_is_number = _generic._is_number
_merged_range_for = _generic._merged_range_for
inspect_workbook = _generic.inspect_workbook


def _section_ranges(sheet: Any, labels: set[str]) -> list[dict[str, Any]]:
    headers = []
    for row in range(1, sheet.max_row + 1):
        for cell in sheet[row]:
            if isinstance(cell.value, str) and cell.value.strip().upper() in labels:
                headers.append({"label": cell.value.strip(), "row": row, "evidence": {**_cell_evidence(cell), "merged_range": _merged_range_for(sheet, cell)}})
    ranges = []
    for index, header in enumerate(headers):
        end_row = headers[index + 1]["row"] if index + 1 < len(headers) else sheet.max_row + 1
        ranges.append({"section_id": f"section-{index + 1}", "label": header["label"], "header_row": header["row"], "end_row_exclusive": end_row, "header": header})
    return ranges


def _empty_row_record(section: dict[str, Any], row_number: int, numeric: list[Any], reason: str) -> dict[str, Any]:
    return {
        "geometric_row_id": f"worksheet-row-{row_number}",
        "worksheet_row": row_number,
        "official_label": None,
        "section_id": section["section_id"],
        "seat_count": 0,
        "unknown_count": 0,
        "gaps": [],
        "rejected_cells": [{"reason": reason, "evidence": _cell_evidence(cell)} for cell in numeric],
        "seats": [],
    }


def _detect_row_markers(row_cells: list[Any], profile: dict[str, Any]):
    numeric = sorted((cell for cell in row_cells if _is_number(cell.value)), key=lambda cell: cell.column)
    strategy = profile["row_detection"]["type"]
    if strategy == "single_letter_markers":
        markers = [cell for cell in row_cells if isinstance(cell.value, str) and len(cell.value.strip()) == 1 and cell.value.strip().isalpha() and cell.value.strip().upper() in profile["row_detection"]["allowed_labels"]]
        if not markers:
            return None, numeric, "no_single_letter_row_marker"
        return {"label": markers[0].value.strip().upper(), "cells": markers}, numeric, None
    if strategy == "equal_numeric_edge_markers":
        if len(numeric) < 2:
            return None, numeric, "insufficient_numeric_edge_markers"
        if numeric[0].value != numeric[-1].value:
            return None, numeric, "edge_numeric_markers_disagree"
        return {"label": str(numeric[0].value), "cells": [numeric[0], numeric[-1]]}, numeric, None
    raise ValueError(f"Unsupported row detection strategy: {strategy}")


def _candidate_cells(row_cells: list[Any], markers: dict[str, Any], profile: dict[str, Any]) -> tuple[list[Any], list[Any], list[dict[str, Any]]]:
    strategy = profile["seat_detection"]["type"]
    numeric = sorted((cell for cell in row_cells if _is_number(cell.value)), key=lambda cell: cell.column)
    rejected = []
    if strategy == "style_allowlist":
        allowed = set(profile["seat_detection"]["style_ids"])
        seats = [cell for cell in numeric if cell.style_id in allowed]
        rejected.extend({"reason": "numeric_service_or_marker", "evidence": _cell_evidence(cell)} for cell in numeric if cell.style_id not in allowed)
        return seats, [], rejected
    if strategy == "numeric_between_edge_markers":
        left, right = markers["cells"][0], markers["cells"][-1]
        interior = [cell for cell in row_cells if left.column < cell.column < right.column]
        seats = [cell for cell in interior if _is_number(cell.value)]
        gaps = [cell for cell in interior if cell.value is None]
        for cell in interior:
            if cell.value is not None and not _is_number(cell.value):
                rejected.append({"reason": "non_numeric_interior_cell", "evidence": _cell_evidence(cell)})
        return seats, gaps, rejected
    raise ValueError(f"Unsupported seat detection strategy: {strategy}")


def extract_full(path: Path, profile: dict[str, Any]) -> dict[str, Any]:
    inspection = inspect_workbook(path)
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheet_name = profile["sheet_name"]
    if sheet_name not in workbook.sheetnames:
        return {"schema": "seating-json-v1-phase8c", "source_only": True, "reference_svg_used": False, "profile_id": profile["profile_id"], "source": inspection, "sections": [], "seats": [], "validation": {"status": "QUARANTINE", "reason": "configured_sheet_missing"}}
    sheet = workbook[sheet_name]
    sections = _section_ranges(sheet, set(profile.get("section_header_labels", [])))
    if not sections:
        sections = [{
            "section_id": "section-1",
            "label": None,
            "header_row": sheet.min_row - 1,
            "end_row_exclusive": sheet.max_row + 1,
            "header": None,
        }]
    rows = []
    seats = []
    gaps = []
    rejected = []
    for section in sections:
        if section["header"] is not None:
            rejected.append({
                "reason": "section_header_not_seat",
                "section_id": section["section_id"],
                "evidence": section["header"]["evidence"],
            })
        for row_number in range(section["header_row"] + 1, section["end_row_exclusive"]):
            row_cells = list(sheet[row_number])
            markers, numeric, marker_error = _detect_row_markers(row_cells, profile)
            if markers is None:
                if numeric:
                    rows.append(_empty_row_record(section, row_number, numeric, marker_error or "row_marker_unknown"))
                    rejected.extend(rows[-1]["rejected_cells"])
                continue
            candidate_cells, row_gaps, row_rejected = _candidate_cells(row_cells, markers, profile)
            row_rejected.extend({
                "reason": "row_label_marker",
                "evidence": _cell_evidence(marker),
            } for marker in markers["cells"])
            row_seats = []
            for cell in candidate_cells:
                number_status = "explicitly_read" if float(cell.value).is_integer() else "unknown"
                seat = {
                    "seat_key": f"xlsx-source-{sheet.title}-{cell.coordinate}",
                    "source_coordinate": _cell_evidence(cell),
                    "computed_geometry": {"grid_x": cell.column, "grid_y": cell.row, "column_letter": cell.column_letter if hasattr(cell, "column_letter") else cell.column},
                    "section_id": section["section_id"],
                    "section_label": section["label"],
                    "geometric_row_id": f"worksheet-row-{row_number}",
                    "row_label": str(markers["label"]),
                    "seat_number": int(cell.value) if number_status == "explicitly_read" else None,
                    "number_status": number_status,
                    "confidence": 1.0 if number_status == "explicitly_read" else 0.0,
                    "evidence": {
                        "sheet": sheet.title,
                        "source_cell": cell.coordinate,
                        "raw_value": cell.value,
                        "style_id": cell.style_id,
                        "row_marker_cells": [marker.coordinate for marker in markers["cells"]],
                        "position_numbering_used": False,
                        "rule_reconstruction_used": False,
                    },
                }
                if number_status == "unknown":
                    row_rejected.append({"reason": "non_integer_numeric_value", "evidence": _cell_evidence(cell)})
                else:
                    row_seats.append(seat)
            row_gap_records = []
            for cell in row_gaps:
                record = {"cell": cell.coordinate, "row": cell.row, "column": cell.column, "reason": "empty_position_inside_row_span", "merged_range": _merged_range_for(sheet, cell)}
                row_gap_records.append(record)
                gaps.append({"section_id": section["section_id"], "geometric_row_id": f"worksheet-row-{row_number}", **record})
            row_record = {
                "geometric_row_id": f"worksheet-row-{row_number}",
                "worksheet_row": row_number,
                "official_label": str(markers["label"]),
                "marker_cells": [marker.coordinate for marker in markers["cells"]],
                "section_id": section["section_id"],
                "section_label": section["label"],
                "seat_count": len(row_seats),
                "unknown_count": sum(seat["number_status"] == "unknown" for seat in row_seats),
                "gaps": row_gap_records,
                "rejected_cells": row_rejected,
                "seats": row_seats,
            }
            rows.append(row_record)
            seats.extend(row_seats)
            rejected.extend(row_rejected)

    identities = [(seat["section_label"], seat["row_label"], seat["seat_number"]) for seat in seats]
    duplicate_identity = sorted({identity for identity in identities if identities.count(identity) > 1})
    duplicate_keys = sorted({seat["seat_key"] for seat in seats if [item["seat_key"] for item in seats].count(seat["seat_key"]) > 1})
    unexpected = [item for item in rejected if item["reason"] not in {"numeric_service_or_marker", "row_label_marker", "section_header_not_seat"}]
    section_summary = []
    for section in sections:
        section_rows = [row for row in rows if row["section_id"] == section["section_id"]]
        section_summary.append({
            "section_id": section["section_id"],
            "label": section["label"],
            "header_row": section["header_row"],
            "row_count": len(section_rows),
            "seat_count": sum(row["seat_count"] for row in section_rows),
            "unknown_count": sum(row["unknown_count"] for row in section_rows),
            "gap_count": sum(len(row["gaps"]) for row in section_rows),
            "rejected_count": sum(len(row["rejected_cells"]) for row in section_rows) + (1 if section["header"] is not None else 0),
        })
    status = "PASS" if not duplicate_identity and not duplicate_keys and not unexpected and all(seat["number_status"] == "explicitly_read" for seat in seats) else "QUARANTINE"
    return {
        "schema": "seating-json-v1-phase8c-full-source-only",
        "source_only": True,
        "reference_svg_used": False,
        "profile_id": profile["profile_id"],
        "source": inspection,
        "sections": section_summary,
        "rows": rows,
        "seats": seats,
        "gaps": gaps,
        "rejected_cells": rejected,
        "summary": {
            "section_count": len(section_summary),
            "row_count": len(rows),
            "seat_count": len(seats),
            "explicit_number_count": sum(seat["number_status"] == "explicitly_read" for seat in seats),
            "unknown_count": sum(seat["number_status"] == "unknown" for seat in seats),
            "gap_count": len(gaps),
            "rejected_count": len(rejected),
            "duplicate_identity_keys": [list(item) for item in duplicate_identity],
            "duplicate_seat_keys": duplicate_keys,
            "position_numbering_used": False,
        },
        "validation": {"status": status, "unexpected_rejection_count": len(unexpected), "all_coordinates_present": all(bool(seat["source_coordinate"]["cell"]) for seat in seats)},
    }

"""Full source-only Hunter College XLSX extraction for PHASE 7B."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "ChatExport_2026-10-11" / "files" / "Площадка .xlsx"
BASE = Path(__file__).resolve().parent
OUT = BASE / "output"
OUT.mkdir(parents=True, exist_ok=True)
SEAT_STYLES = {4, 6}
ROW_RE = re.compile(r"^[A-Z]$")


def coordinate(cell, min_row, max_row, min_col, max_col):
    return {
        "sheet": cell.parent.title,
        "cell": cell.coordinate,
        "row": cell.row,
        "column": cell.column,
        "column_letter": get_column_letter(cell.column),
        "grid_x": cell.column,
        "grid_y": cell.row,
        "grid_x_normalized": round((cell.column - min_col) / (max_col - min_col), 6),
        "grid_y_normalized": round((cell.row - min_row) / (max_row - min_row), 6),
    }


def main():
    workbook = load_workbook(SOURCE, data_only=False, read_only=False)
    sheet = workbook["Scheme"]
    used_range = sheet.calculate_dimension()
    min_col, min_row, max_col, max_row = sheet.min_column, sheet.min_row, sheet.max_column, sheet.max_row

    # Read rows A-Z from the actual row markers.  Repeated markers are kept as
    # evidence; only the canonical leftmost marker names the logical row.
    logical_rows = []
    for worksheet_row in range(min_row, max_row + 1):
        cells = list(sheet[worksheet_row])
        markers = [cell for cell in cells if isinstance(cell.value, str) and ROW_RE.fullmatch(cell.value)]
        if not markers:
            continue
        label = markers[0].value
        if label not in "ABCDEFGHIJKLMNOPQRSTUVWXYZ":
            continue
        numeric = [cell for cell in cells if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)]
        seats = [cell for cell in numeric if cell.style_id in SEAT_STYLES]
        excluded_numeric = [cell for cell in numeric if cell.style_id not in SEAT_STYLES]
        if not seats:
            continue
        logical_rows.append({
            "label": label,
            "worksheet_row": worksheet_row,
            "markers": markers,
            "numeric": numeric,
            "seats": sorted(seats, key=lambda cell: cell.column),
            "excluded_numeric": excluded_numeric,
        })

    logical_rows.sort(key=lambda row: row["worksheet_row"])
    if [row["label"] for row in logical_rows] != list("ABCDEFGHIJKLMNOPQRSTUVWXYZ"):
        raise RuntimeError(f"Expected A-Z rows, found {[row['label'] for row in logical_rows]}")

    all_seats = []
    rows = []
    for row in logical_rows:
        seat_cells = row["seats"]
        blocks = []
        for cell in seat_cells:
            if not blocks or cell.column != blocks[-1]["last_column"] + 1:
                blocks.append({"first_column": cell.column, "last_column": cell.column, "cells": []})
            else:
                blocks[-1]["last_column"] = cell.column
            blocks[-1]["cells"].append(cell)

        row_seats = []
        for block_index, block in enumerate(blocks, start=1):
            for position, cell in enumerate(block["cells"], start=1):
                seat = {
                    "seat_key": f"hunter-source-row-{row['label']}-cell-{cell.coordinate}",
                    "source_coordinate": coordinate(cell, min_row, max_row, min_col, max_col),
                    "computed_geometry": {
                        "grid_x": cell.column,
                        "grid_y": cell.row,
                        "relative_x": round(cell.column - min_col, 6),
                        "relative_y": round(cell.row - min_row, 6),
                        "block_index": block_index,
                        "position_in_block": position,
                        "block_size": len(block["cells"]),
                    },
                    "row_label": row["label"],
                    "seat_number": int(cell.value),
                    "number_status": "explicitly_read",
                    "confidence": 1.0,
                    "evidence": {
                        "cell_value": cell.value,
                        "cell_value_type": cell.data_type,
                        "seat_style_id": cell.style_id,
                        "fill_rgb": cell.fill.fgColor.rgb,
                        "source_cell": cell.coordinate,
                        "position_numbering_used": False,
                        "rule_reconstruction_used": False,
                    },
                }
                row_seats.append(seat)
                all_seats.append(seat)

        number_values = [seat["seat_number"] for seat in row_seats]
        duplicate_numbers = sorted({value for value in number_values if number_values.count(value) > 1})
        rows.append({
            "row_label": row["label"],
            "worksheet_row": row["worksheet_row"],
            "row_label_cells": [cell.coordinate for cell in row["markers"]],
            "seat_count": len(row_seats),
            "numeric_cell_count": len(row["numeric"]),
            "excluded_numeric_cells": [cell.coordinate for cell in row["excluded_numeric"]],
            "blocks": [
                {
                    "block_index": index,
                    "size": len(block["cells"]),
                    "first_column": get_column_letter(block["first_column"]),
                    "last_column": get_column_letter(block["last_column"]),
                }
                for index, block in enumerate(blocks, start=1)
            ],
            "number_values_observed": number_values,
            "duplicate_numbers_within_row": duplicate_numbers,
            "explicitly_read_count": len(row_seats),
            "rule_reconstructed_count": 0,
            "unknown_count": 0,
            "seats": row_seats,
        })

    key_pairs = [(seat["row_label"], seat["seat_number"]) for seat in all_seats]
    duplicate_keys = sorted({pair for pair in key_pairs if key_pairs.count(pair) > 1})
    result = {
        "schema": "hunter-phase7b-source-only-v1",
        "source_only": True,
        "reference_svg_used": False,
        "source": {
            "path": str(SOURCE),
            "sha256": hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
            "workbook_sheets": workbook.sheetnames,
            "sheet": sheet.title,
            "used_range": used_range,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "merged_cells": [str(rng) for rng in sheet.merged_cells.ranges],
            "formula_cells": [cell.coordinate for row_cells in sheet.iter_rows() for cell in row_cells if cell.data_type == "f"],
            "hidden_rows": [idx for idx, dim in sheet.row_dimensions.items() if dim.hidden],
            "hidden_columns": [idx for idx, dim in sheet.column_dimensions.items() if dim.hidden],
            "global_numeric_cell_count": sum(
                isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
                for row_cells in sheet.iter_rows()
                for cell in row_cells
            ),
            "global_seat_style_cell_count": sum(
                isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) and cell.style_id in SEAT_STYLES
                for row_cells in sheet.iter_rows()
                for cell in row_cells
            ),
        },
        "rows": rows,
        "seats": all_seats,
        "summary": {
            "row_count": len(rows),
            "seat_count": len(all_seats),
            "explicitly_read_count": sum(seat["number_status"] == "explicitly_read" for seat in all_seats),
            "rule_reconstructed_count": 0,
            "unknown_count": 0,
            "duplicate_row_seat_keys": [list(pair) for pair in duplicate_keys],
            "duplicate_numbers_within_any_row": [row["row_label"] for row in rows if row["duplicate_numbers_within_row"]],
            "all_coordinates_in_used_range": all(
                min_col <= seat["source_coordinate"]["column"] <= max_col
                and min_row <= seat["source_coordinate"]["row"] <= max_row
                for seat in all_seats
            ),
            "position_numbering_used": False,
        },
    }
    output = OUT / "seating-model-source-only.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "rows": len(rows), "seats": len(all_seats)}, ensure_ascii=False))


if __name__ == "__main__":
    main()

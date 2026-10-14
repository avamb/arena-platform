"""Independent evaluator for the bounded Aribau transfer test."""

from __future__ import annotations

import hashlib
import json
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from openpyxl import load_workbook


BASE = Path(__file__).resolve().parent
ROOT = BASE.parents[1]
SOURCE_JSON = BASE / "output" / "aribau-row21-source-only.json"
SOURCE_XLSX = ROOT / "ChatExport_2026-10-11" / "files" / "AFORO ARIBAU SALA Barcelona.xlsx"
REFERENCE = ROOT / "ChatExport_2026-10-11" / "files" / "aforo_aribau.svg"
SVG_NS = "{http://www.w3.org/2000/svg}"
SECTION_LABELS = {"BALCONY", "PARTERRE", "STAGE"}


def label_of(element):
    return next((value for key, value in element.attrib.items() if key.endswith("}label") or key == "inkscape:label"), None)


def section_source_facts():
    workbook = load_workbook(SOURCE_XLSX, data_only=False, read_only=False)
    sheet = workbook["Aribau,Sala 5 Aforo ok"]
    header_row = next(row for row in range(1, sheet.max_row + 1) if any(isinstance(cell.value, str) and cell.value.strip().upper() == "PARTERRE" for cell in sheet[row]))
    next_header = next((row for row in range(header_row + 1, sheet.max_row + 1) if any(isinstance(cell.value, str) and cell.value.strip().upper() == "STAGE" for cell in sheet[row])), sheet.max_row + 1)
    rows = []
    for row_number in range(header_row + 1, next_header):
        numeric = [cell for cell in sheet[row_number] if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)]
        numeric.sort(key=lambda cell: cell.column)
        if len(numeric) < 3 or numeric[0].value != numeric[-1].value:
            continue
        rows.append({"worksheet_row": row_number, "row_label": str(numeric[0].value), "seat_count": len(numeric) - 2})
    return {"section_label": "PARTERRE", "header_row": header_row, "row_count": len(rows), "seat_count": sum(row["seat_count"] for row in rows), "rows": rows}


def reference_facts():
    root = ET.parse(REFERENCE).getroot()
    groups = []
    numeric_titles = 0
    for group in root.iter(SVG_NS + "g"):
        if label_of(group) not in {"#Parterre", "#Parter"}:
            continue
        circles = list(group.iter(SVG_NS + "circle"))
        titles = []
        for circle in circles:
            title = next(("".join(node.itertext()).strip() for node in circle.iter() if node.tag.split("}")[-1] == "title"), None)
            titles.append(title)
            if title is not None and re.fullmatch(r"\d{1,4}", title):
                numeric_titles += 1
        groups.append({"label": label_of(group), "circle_count": len(circles), "titles": titles})
    return {
        "path": str(REFERENCE),
        "sha256": hashlib.sha256(REFERENCE.read_bytes()).hexdigest(),
        "group_count": len(groups),
        "seat_count": sum(group["circle_count"] for group in groups),
        "group_sizes": [group["circle_count"] for group in groups],
        "numeric_title_count": numeric_titles,
    }


def main():
    source_only = json.loads(SOURCE_JSON.read_text(encoding="utf-8"))
    if not source_only["source_only"] or source_only["reference_svg_used"]:
        raise RuntimeError("Source-only latch is invalid")
    source = section_source_facts()
    reference = reference_facts()
    mismatch_reasons = []
    if source["seat_count"] != reference["seat_count"]:
        mismatch_reasons.append(f"PARTERRE source has {source['seat_count']} seats; SVG #Parterre has {reference['seat_count']}")
    if source["row_count"] != reference["group_count"]:
        mismatch_reasons.append(f"PARTERRE source has {source['row_count']} rows; SVG has {reference['group_count']} #Parterre groups")
    version_status = "VERSION_MISMATCH" if mismatch_reasons else "STRUCTURALLY_COMPATIBLE"
    result = {
        "schema": "aribau-phase8b-independent-evaluation-v1",
        "source_only_result": str(SOURCE_JSON),
        "reference": reference,
        "version_check": {
            "source": source,
            "status": version_status,
            "mismatch_reasons": mismatch_reasons,
            "row21_source_seat_count": len(source_only["seats"]),
            "reference_group_size_matches_row21": all(size == len(source_only["seats"]) for size in reference["group_sizes"]),
        },
        "metrics": {
            "geometry_precision": None,
            "geometry_recall": None,
            "number_precision": None,
            "number_recall": None,
            "exact_seat_identity_accuracy": None,
            "status": "unknown_due_version_mismatch" if mismatch_reasons else "not_run_bounded_row_only",
        },
        "policy": {
            "reference_used_only_after_source_latch": True,
            "extractor_changed_after_reference_read": False,
            "accuracy_metrics_confirmed": False,
        },
    }
    output = BASE / "output" / "aribau-independent-evaluation.json"
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(output), "version_status": version_status, "source_partierre_seats": source["seat_count"], "reference_parterre_seats": reference["seat_count"], "metrics_confirmed": False}, ensure_ascii=False))


if __name__ == "__main__":
    main()

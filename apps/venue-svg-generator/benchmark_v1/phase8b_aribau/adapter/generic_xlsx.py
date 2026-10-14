"""Small fail-closed XLSX adapter used by the Phase 8B transfer test."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def _is_number(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def inspect_workbook(path: Path) -> dict[str, Any]:
    """Read workbook structure without classifying any cell as a seat."""
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        sheets.append({
            "title": sheet.title,
            "used_range": sheet.calculate_dimension(),
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "merged_cells": [str(item) for item in sheet.merged_cells.ranges],
            "formula_cells": [cell.coordinate for row in sheet.iter_rows() for cell in row if cell.data_type == "f"],
            "hidden_rows": [idx for idx, dim in sheet.row_dimensions.items() if dim.hidden],
            "hidden_columns": [idx for idx, dim in sheet.column_dimensions.items() if dim.hidden],
            "nonempty_cell_count": sum(cell.value is not None for row in sheet.iter_rows() for cell in row),
            "numeric_cell_count": sum(_is_number(cell.value) for row in sheet.iter_rows() for cell in row),
        })
    return {
        "path": str(path),
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "workbook_sheets": workbook.sheetnames,
        "sheets": sheets,
    }


def _cell_evidence(cell: Any) -> dict[str, Any]:
    fill = cell.fill.fgColor
    return {
        "sheet": cell.parent.title,
        "cell": cell.coordinate,
        "row": cell.row,
        "column": cell.column,
        "column_letter": get_column_letter(cell.column),
        "raw_value": cell.value,
        "data_type": cell.data_type,
        "style_id": cell.style_id,
        "fill_rgb": str(getattr(fill, "rgb", fill)) if fill is not None else None,
    }


def _merged_range_for(sheet: Any, cell: Any) -> str | None:
    for merged in sheet.merged_cells.ranges:
        if cell.coordinate in merged:
            return str(merged)
    return None


def _nearest_section_header(sheet: Any, row_number: int, known_labels: set[str]) -> dict[str, Any] | None:
    """Find the nearest textual section header above a row, including merged anchors."""
    for candidate_row in range(row_number - 1, 0, -1):
        for cell in sheet[candidate_row]:
            if isinstance(cell.value, str) and cell.value.strip().upper() in known_labels:
                return {
                    "label": cell.value.strip(),
                    "anchor": cell.coordinate,
                    "merged_range": _merged_range_for(sheet, cell),
                    "evidence": _cell_evidence(cell),
                }
    return None


def _quarantine(inspection: dict[str, Any], profile: dict[str, Any], reason: str) -> dict[str, Any]:
    return {
        "schema": "seating-json-v1-phase8b-aribau-source-only",
        "source_only": True,
        "reference_svg_used": False,
        "profile_id": profile["profile_id"],
        "source": inspection,
        "section": None,
        "row": None,
        "seats": [],
        "gaps": [],
        "rejected_cells": [{"reason": reason}],
        "validation": {
            "seat_count": 0,
            "explicit_number_count": 0,
            "unknown_number_count": 0,
            "duplicate_seat_keys": [],
            "duplicate_identity_keys": [],
            "all_seat_coordinates_present": True,
            "position_numbering_used": False,
            "status": "QUARANTINE",
        },
    }


def extract_row_source_only(path: Path, profile: dict[str, Any]) -> dict[str, Any]:
    """Extract one configured row using structural rules and no reference data."""
    inspection = inspect_workbook(path)
    workbook = load_workbook(path, data_only=False, read_only=False)
    sheet_name = profile["sheet_name"]
    if sheet_name not in workbook.sheetnames:
        return _quarantine(inspection, profile, f"configured_sheet_missing:{sheet_name}")
    sheet = workbook[sheet_name]
    row_number = int(profile["row_scope"]["worksheet_row"])
    row_cells = list(sheet[row_number])
    section = _nearest_section_header(sheet, row_number, set(profile["section_header_labels"]))
    if section is None:
        return _quarantine(inspection, profile, "section_header_unknown")

    numeric = [cell for cell in row_cells if _is_number(cell.value)]
    if len(numeric) < 3:
        return _quarantine(inspection, profile, "insufficient_numeric_cells_for_row_markers")
    numeric.sort(key=lambda cell: cell.column)
    left_marker, right_marker = numeric[0], numeric[-1]
    if left_marker.value != right_marker.value:
        return _quarantine(inspection, profile, "edge_row_markers_disagree")
    if left_marker.column >= right_marker.column - 1:
        return _quarantine(inspection, profile, "empty_interior_after_row_markers")

    marker_rejected = []
    for marker, reason in ((left_marker, "row_label_marker"), (right_marker, "row_label_marker")):
        marker_rejected.append({"reason": reason, "evidence": _cell_evidence(marker)})
    merged_rejected = []
    for merged in sheet.merged_cells.ranges:
        if merged.min_row <= row_number - 1 <= merged.max_row and merged.max_row < row_number:
            anchor = sheet.cell(merged.min_row, merged.min_col)
            if isinstance(anchor.value, str) and anchor.value.strip().upper() in set(profile["section_header_labels"]):
                merged_rejected.append({
                    "reason": "section_header_not_seat",
                    "evidence": {**_cell_evidence(anchor), "merged_range": str(merged)},
                })

    interior = [cell for cell in row_cells if left_marker.column < cell.column < right_marker.column]
    seats = []
    gaps = []
    rejected = marker_rejected + merged_rejected
    for cell in interior:
        evidence = _cell_evidence(cell)
        if cell.value is None:
            gaps.append({"cell": cell.coordinate, "row": cell.row, "column": cell.column, "reason": "empty_position_within_row_span"})
        elif _is_number(cell.value):
            seats.append({
                "seat_key": f"xlsx-source-{sheet.title}-{cell.coordinate}",
                "source_coordinate": evidence,
                "row_label": str(left_marker.value),
                "section_label": section["label"],
                "seat_number": int(cell.value) if float(cell.value).is_integer() else None,
                "number_status": "explicitly_read" if float(cell.value).is_integer() else "unknown",
                "confidence": 1.0 if float(cell.value).is_integer() else 0.0,
                "evidence": {
                    "source_cell": cell.coordinate,
                    "raw_value": cell.value,
                    "style_id": cell.style_id,
                    "row_marker_cells": [left_marker.coordinate, right_marker.coordinate],
                    "position_numbering_used": False,
                    "rule_reconstruction_used": False,
                },
            })
            if not float(cell.value).is_integer():
                rejected.append({"reason": "non_integer_numeric_value", "evidence": evidence})
                seats.pop()
        else:
            rejected.append({"reason": "non_numeric_interior_cell", "evidence": evidence})

    duplicate_keys = sorted({seat["seat_key"] for seat in seats if [item["seat_key"] for item in seats].count(seat["seat_key"]) > 1})
    duplicate_identity = sorted({(seat["section_label"], seat["row_label"], seat["seat_number"]) for seat in seats if [(item["section_label"], item["row_label"], item["seat_number"]) for item in seats].count((seat["section_label"], seat["row_label"], seat["seat_number"])) > 1})
    unknown = [seat for seat in seats if seat["number_status"] == "unknown"]
    unexpected_rejections = [item for item in rejected if item["reason"] not in {"row_label_marker", "section_header_not_seat"}]
    result = {
        "schema": "seating-json-v1-phase8b-aribau-source-only",
        "source_only": True,
        "reference_svg_used": False,
        "profile_id": profile["profile_id"],
        "source": inspection,
        "section": section,
        "row": {
            "worksheet_row": row_number,
            "geometric_row_id": f"worksheet-row-{row_number}",
            "official_label": str(left_marker.value),
            "marker_cells": [left_marker.coordinate, right_marker.coordinate],
            "marker_detection": "equal_numeric_edge_markers",
        },
        "seats": seats,
        "gaps": gaps,
        "rejected_cells": rejected,
        "validation": {
            "seat_count": len(seats),
            "explicit_number_count": sum(seat["number_status"] == "explicitly_read" for seat in seats),
            "unknown_number_count": len(unknown),
            "duplicate_seat_keys": duplicate_keys,
            "duplicate_identity_keys": [list(item) for item in duplicate_identity],
            "all_seat_coordinates_present": all(bool(seat["source_coordinate"]["cell"]) for seat in seats),
            "position_numbering_used": False,
            "status": "PASS" if section and not unknown and not duplicate_keys and not duplicate_identity and not unexpected_rejections else "QUARANTINE",
        },
    }
    return result


def write_json(result: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
